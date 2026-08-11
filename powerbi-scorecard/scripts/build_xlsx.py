#!/usr/bin/env python3
"""Genera Scorecard_Model.xlsx: libro fuente para Power BI con 3 tablas
(Dim_KPI, Dim_Date, Fact_Scorecard) como Tablas de Excel con nombre."""
import csv, os
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

DATA = "/home/user/angela/powerbi-scorecard/data"
OUT  = os.path.join(DATA, "Scorecard_Model.xlsx")

def read_csv(name):
    with open(os.path.join(DATA, name), encoding="utf-8-sig") as f:
        return list(csv.reader(f))

HEADER_FILL = PatternFill("solid", fgColor="1F3864")   # azul DS Smith oscuro
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT   = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def num_or_str(v):
    if v is None or v == "":
        return None
    try:
        if "." in v or "e" in v.lower():
            return float(v)
        return int(v)
    except (ValueError, AttributeError):
        return v

def add_sheet(wb, title, rows, table_name, date_cols=(), num_cols=(), pct_cols=(), col_widths=None):
    ws = wb.create_sheet(title)
    header, body = rows[0], rows[1:]
    ws.append(header)
    for rrow in body:
        out = []
        for j, val in enumerate(rrow):
            if j in date_cols:
                out.append(date.fromisoformat(val) if val else None)
            elif j in num_cols or j in pct_cols:
                out.append(num_or_str(val))
            else:
                out.append(val if val != "" else None)
        ws.append(out)
    # estilos
    for j, _ in enumerate(header, 1):
        c = ws.cell(row=1, column=j)
        c.fill = HEADER_FILL; c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="left")
        c.border = BORDER
    for r in range(2, ws.max_row + 1):
        for j in range(1, len(header) + 1):
            c = ws.cell(row=r, column=j)
            c.font = BODY_FONT; c.border = BORDER
            if (j - 1) in date_cols:
                c.number_format = "yyyy-mm-dd"
            elif (j - 1) in pct_cols:
                c.number_format = "0.0%"
            elif (j - 1) in num_cols:
                c.number_format = "#,##0.###"
    # anchos
    if col_widths:
        for j, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    # tabla con nombre
    ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"
    tbl = Table(displayName=table_name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)
    return ws

wb = openpyxl.Workbook()
wb.remove(wb.active)

# Portada / índice
info = wb.create_sheet("_Info")
info["B2"] = "DS Smith SSC – Balanced Scorecard | Modelo de datos para Power BI"
info["B2"].font = Font(name="Arial", size=13, bold=True, color="1F3864")
notes = [
    "",
    "Fuente: 2026_AP07_DE_DSS_Balanced_Scorecard_SSC (modelo de grupo, hoja Data_DE_Pack).",
    "Este libro NO sustituye al modelo de grupo: es la capa de datos limpia y en formato",
    "largo (tidy) que consume Power BI.",
    "",
    "Tablas (una por hoja, como Tabla de Excel con nombre):",
    "   • Dim_KPI         – catálogo de KPIs (perspectiva, objetivo, proceso, unidad, target, dirección).",
    "   • Dim_Date        – calendario mensual con año fiscal DS Smith (mayo–abril).",
    "   • Fact_Scorecard  – un registro por KPI y mes: Numerator y Denominator (datos crudos).",
    "",
    "El valor del KPI se calcula en Power BI con medidas DAX (ver carpeta /dax) según Value_Type:",
    "   PCT = Numerator/Denominator ·  RATE1000 = Num/Den×1000 ·  RATIO = Num/Den ·",
    "   COUNT / CURRENCY = Numerator.",
    "",
    "Relaciones: Dim_KPI[KPI_ID] 1—* Fact_Scorecard[KPI_ID]  ·  Dim_Date[Date] 1—* Fact_Scorecard[Date].",
    "",
    "Actualización: pega los nuevos meses de Data_DE_Pack en Fact_Scorecard (o re-ejecuta el",
    "script de extracción) y pulsa Actualizar en Power BI.",
]
for i, t in enumerate(notes, 4):
    c = info.cell(row=i, column=2, value=t)
    c.font = Font(name="Arial", size=10, bold=t.strip().startswith("Tablas"))
info.column_dimensions["A"].width = 2
info.sheet_view.showGridLines = False

add_sheet(wb, "Dim_KPI", read_csv("Dim_KPI.csv"), "Dim_KPI",
          pct_cols=(10,),  # Target (algunos son %)
          col_widths=[8,30,45,20,8,34,55,12,14,10,10,10])
add_sheet(wb, "Dim_Date", read_csv("Dim_Date.csv"), "Dim_Date",
          date_cols=(0,), num_cols=(1,2,6),
          col_widths=[12,8,8,12,10,12,12])
add_sheet(wb, "Fact_Scorecard", read_csv("Fact_Scorecard.csv"), "Fact_Scorecard",
          date_cols=(1,), num_cols=(2,3),
          col_widths=[10,12,14,14])

wb.save(OUT)
print("Escrito", OUT)
